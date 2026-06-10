"""
depot_importer.py — Parses the per-depot Excel workbook (DEPOT_ORDERS_*.xlsx).

Each depot has its own sheet.  Each sheet has two sections:
  Section 1 — depot totals (skipped; we only use Section 2)
  Section 2 — loading breakdown (starts after a row whose col-0 contains '▸')

The breakdown header row pattern:
  ['PRODUCT DESCRIPTION', 'HATCLIFF 1', None, 'HATCLIFF 2', None, 'HATCLIFF 3', ...]
  Truck labels are at odd column indices (1, 3, 5 …).

Returns a flat list of depot_orders rows ready for upsert.
"""

from __future__ import annotations

import io
import logging
from datetime import date
from typing import Union

import openpyxl

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

SKIP_SHEETS = {"Sheet1", "PARAMETERS"}

BREAD_SKUS = {
    "SUPERIOR", "BROWN", "WHOLE GRAIN", "WHOLE WHEAT", "WHOLEGRAIN",
    "MR CHINGWA", "MRS CHINGWA", "DR CHINGWA",
    "SPAR WHITE", "SPAR BROWN", "SPAR WHOLE GRAIN",
}

# Prefixes that mark rows to skip entirely
SKIP_ROW_PREFIXES = (
    "TOTAL", "GRAND TOTAL", "SUPERVISOR", "HEADLOADER",
    "SUPERVISOR OPERATIONAL", "TIME & BAY", "TRUCK REG",
)


def _normalise(val) -> str:
    """Coerce a cell value to a stripped uppercase string."""
    if val is None:
        return ""
    return str(val).strip().upper()


def _int_or_zero(val) -> int:
    if val is None:
        return 0
    try:
        return int(float(val))
    except (ValueError, TypeError):
        return 0


def classify_sku_group(sku_name: str) -> str:
    return "bread" if sku_name.upper() in BREAD_SKUS else "confect"


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def parse_depot_excel(
    file_obj: Union[bytes, io.BytesIO],
    dispatch_date: date,
) -> tuple[list[dict], list[str]]:
    """
    Parse every depot sheet in *file_obj* and return
    (list_of_depot_order_dicts, warnings).
    """
    if isinstance(file_obj, bytes):
        file_obj = io.BytesIO(file_obj)

    warnings: list[str] = []
    all_rows: list[dict] = []

    try:
        wb = openpyxl.load_workbook(file_obj, data_only=True)
    except Exception as exc:
        msg = f"Failed to open workbook: {exc}"
        logger.error(msg)
        return [], [msg]

    for sheet_name in wb.sheetnames:
        if sheet_name.strip() in SKIP_SHEETS:
            continue
        ws = wb[sheet_name]
        depot_name = sheet_name.strip().upper()

        rows_raw = list(ws.iter_rows(values_only=True))

        # ── Find the breakdown section marker ──────────────────────────
        breakdown_idx = None
        for i, row in enumerate(rows_raw):
            first = _normalise(row[0]) if row else ""
            if first.startswith("▸") or "LOADING BREAKDOWN" in first:
                breakdown_idx = i
                break

        if breakdown_idx is None:
            warnings.append(f"Sheet '{sheet_name}': no '▸ LOADING BREAKDOWN' row found — skipped.")
            continue

        section = rows_raw[breakdown_idx + 1:]   # everything after the marker
        if not section:
            warnings.append(f"Sheet '{sheet_name}': breakdown section is empty — skipped.")
            continue

        # ── Row 0 of section: truck header ─────────────────────────────
        header_row = section[0]
        truck_labels: list[str] = []
        truck_col_indices: list[int] = []
        for col_idx in range(1, len(header_row)):
            val = _normalise(header_row[col_idx])
            if val and val not in ("", "NONE"):
                truck_labels.append(val.title())   # e.g. "Hatcliff 1" → title case
                truck_col_indices.append(col_idx)

        if not truck_labels:
            warnings.append(f"Sheet '{sheet_name}': could not parse truck header — skipped.")
            continue

        # ── Row 1: TRUCK REG ───────────────────────────────────────────
        truck_registrations: list[str] = [""] * len(truck_labels)
        if len(section) > 1:
            reg_row = section[1]
            first = _normalise(reg_row[0])
            if first == "TRUCK REG" or first == "":
                for t_idx, col_idx in enumerate(truck_col_indices):
                    if col_idx < len(reg_row):
                        reg = _normalise(reg_row[col_idx])
                        truck_registrations[t_idx] = reg if reg else ""

        # ── Remaining rows: SKU data ───────────────────────────────────
        # Start from row 2 (skip TRUCK REG); also skip TIME & BAY
        sku_start = 2
        for row in section[sku_start:]:
            first = _normalise(row[0])
            if not first:
                continue
            # Stop at operational notes
            if any(first.startswith(p) for p in ("SUPERVISOR", "HEADLOADER")):
                break
            # Skip total / header rows
            if any(first.startswith(p) for p in SKIP_ROW_PREFIXES):
                continue

            sku_name = row[0]
            if sku_name is None:
                continue
            sku_name = str(sku_name).strip()
            if not sku_name:
                continue

            sku_group = classify_sku_group(sku_name)

            for t_idx, col_idx in enumerate(truck_col_indices):
                qty = _int_or_zero(row[col_idx] if col_idx < len(row) else None)

                all_rows.append({
                    "dispatch_date":      dispatch_date.isoformat(),
                    "depot_name":         depot_name,
                    "truck_label":        truck_labels[t_idx],
                    "truck_registration": truck_registrations[t_idx],
                    "sku_name":           sku_name,
                    "sku_group":          sku_group,
                    "ordered_qty":        qty,
                    "loaded_qty":         0,
                    "status":             "In Queue",
                })

        logger.info(
            "parse_depot_excel: sheet '%s' → %d trucks, rows so far %d",
            sheet_name, len(truck_labels), len(all_rows),
        )

    return all_rows, warnings
